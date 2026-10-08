# -*- coding: utf-8 -*-
"""Строит seed/data.json из таблицы «Шаблон деталей.xlsx» — путь детали по процессам АРМ.

Таблица: шапка — колонки матрицы с кодами (А1…Т2), ниже — блоки-сценарии. Каждая строка блока — состояние одной и той же
строки заказа на очередном шаге процесса (колонка Н5 «Статус заявки»), Т1 — название сценария, Т2 — узел/пояснение шага.
Шапка повторяется перед каждым блоком; строки без Н5 (легенда цветов, примечания) пропускаются.

seed создаёт по заказу покупателя на сценарий (номер клиента SEED-<В5>-<n>) и по строке АРМ на каждый шаг —
в АРМ виден весь путь детали, строка за строкой, как в таблице.

    python tools/seed_from_xlsx.py                       # seed/Шаблон деталей.xlsx -> seed/data.json
    python tools/seed_from_xlsx.py <xlsx> <data.json>

Нужен openpyxl (pip install openpyxl) — только для пересборки data.json; seed-1c.ps1 его не требует.
"""
import datetime
import io
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, 'seed', 'Шаблон деталей.xlsx')
OUT = os.path.join(ROOT, 'seed', 'data.json')

# Шаг процесса (Н5 таблицы, без регистра) -> значение Арм_СтатусыАРМ.
# «Заказ поставщику» с О5 «Оплачено» = да — уже «Оплачено» (см. step_status).
STEP_STATUS = {
    'новая': 'Черновик',
    'в работе': 'ВРаботе',
    'заказ поставщику': 'Заказано',
    'приемка упд': 'Приходуется',
    'м3 резер автоматом пропустили': 'Резервирование',
    'резерв': 'Резервирование',
    'размещение по складам': 'НаСкладе',
    'уведомление о приходе': 'НаСкладе',
    'задание на комплектацию': 'Комплектуется',
    'комплектация сборочный лист': 'Комплектуется',
    'контроль комплектации': 'Собран',
    'уведомление о готовности': 'Подготовлено',
    'отгрузочная ведомость': 'Подготовлено',
    'отгрузка клиенту': 'Отгружено',
    'выдача в ремзону': 'Отгружено',
    'возрат товра': 'Возврат',
    'приёмка / подготовка возврата на складе': 'Возврат',
    'контоль документов': 'Выполнено',
    'контроль документов': 'Выполнено',
    'печать комплекта документов и выдача авто': 'Завершено',
}
# статусы, на которых строка уже отправлена на сборку (В9 «Дата отправки на сборку», v2.14)
AFTER_ASSEMBLY = {'Комплектуется', 'Собран', 'Подготовлено', 'Отгружено', 'Выполнено', 'Завершено'}

PAYMENT_FORMS = {'0.22': 'Безнал (НДС 22%)', 'нал': 'Наличные'}   # О3 таблицы -> Арм_ФормыОплаты
PAID = {'да': 'Оплачено'}                                          # О5 таблицы -> поле «Оплачено»

# код колонки -> поле регистра Арм_ДанныеЗакупокИПродаж (generate_arm_matrix_rights._FIELD_PAIRS)
FIELDS = {
    'В1': 'ДатаОтгрузки', 'В2': 'НомерОтгрузки', 'В3': 'ДатаЗаказа', 'В8': 'ВидУслуги',
    'Е1': 'Марка', 'Е2': 'Модель', 'Е3': 'Агрегат', 'Е4': 'Госномер', 'Е5': 'НомерШасси', 'Е6': 'VIN', 'Е7': 'Двигатель',
    'М1': 'Номенклатура', 'Н0': 'Количество', 'Н1': 'ОстатокДляСтроки', 'Н2': 'КоличествоВРезерве', 'Н3': 'КоличествоВПути',
    'Н6': 'ЯчейкаСклада', 'О1': 'СебестоимостьЕдиницы', 'О2': 'Себестоимость', 'О3': 'ФормаОплаты', 'О4': 'РРЦ',
    'О5': 'Оплачено', 'О6': 'НомерСчета', 'О7': 'ДатаСчета', 'Р0': 'КоличествоНормаЧас', 'Р1': 'Коэффициент',
    'Р2': 'Цена', 'С1': 'Поставщик', 'С2': 'ДатаПоступления', 'С3': 'ДатаУПД', 'С4': 'НомерУПД',
    'С5': 'ТерриторияОтгрузкиПоставщиком', 'С6': 'ДатаЗаказаПоставщику', 'С7': 'СрокПоставки', 'Т1': 'КомментарийКСтроке',
}
NUMBER_FIELDS = {'Количество', 'ОстатокДляСтроки', 'КоличествоВРезерве', 'КоличествоВПути', 'СебестоимостьЕдиницы',
                 'Себестоимость', 'РРЦ', 'КоличествоНормаЧас', 'Коэффициент', 'Цена'}
# код колонки -> реквизит строки «Запасы» заказа покупателя (поля «ФАКТ» — как прислал клиент)
CLIENT = {'К1': 'СрочностьКлиента', 'К2': 'ГРЗКлиента', 'К3': 'ТерриторияКлиента', 'К4': 'НоменклатураКлиента',
          'К5': 'АртикулКлиента', 'К6': 'КоличествоКлиента', 'К7': 'ЕдиницаИзмеренияКлиента'}


def blank(v):
    return v is None or (isinstance(v, str) and (not v.strip() or set(v.strip()) <= {'-'}))   # «------» — нет значения


def number(v):
    if not isinstance(v, (int, float)):
        m = re.match(r'\s*(-?\d+(?:[.,]\d+)?)', str(v))    # «200/182» (в пути / свободно) -> 200
        if not m:
            return None
        v = float(m.group(1).replace(',', '.'))
    v = round(float(v), 4)
    return int(v) if v.is_integer() else v


def text(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def date(v):
    return v.strftime('%Y-%m-%d') if isinstance(v, (datetime.date, datetime.datetime)) else text(v)


def evaluate(formula, row_cells):
    """=AE11*AK11 — значения ячеек той же строки; пустая ячейка = 0."""
    def ref(m):
        n = number(row_cells.get(m.group(1)))
        return repr(n or 0)
    expr = re.sub(r'\$?([A-Z]{1,3})\$?\d+', ref, formula.lstrip('='))
    if not re.fullmatch(r'[\d.+\-*/() ]+', expr):
        raise ValueError('формула не поддержана: %s' % formula)
    return number(eval(expr))    # noqa: S307 — только числа и арифметика (проверено выше)


def step_status(step, paid):
    s = STEP_STATUS.get(step.strip().lower())
    if s is None:
        raise ValueError('шаг «%s» не сопоставлен статусу АРМ — дополните STEP_STATUS' % step)
    if s == 'Заказано' and paid:
        s = 'Оплачено'
    return s


def build(xlsx):
    ws = openpyxl.load_workbook(xlsx).worksheets[0]
    code_by_col = {}
    scenarios, cur = [], None
    nomenclature, counterparties, service_types, payment_forms = {}, {}, [], []

    for r in range(1, ws.max_row + 1):
        cells = {c.column_letter: c.value for c in ws[r] if c.value is not None}
        heads = {col: re.match(r'^([А-ЯЁ]\d)', str(v)) for col, v in cells.items()}
        if sum(1 for m in heads.values() if m) > 20:                 # строка шапки
            code_by_col = {col: m.group(1) for col, m in heads.items() if m}
            continue
        v = {code_by_col[col]: val for col, val in cells.items() if col in code_by_col}
        if not code_by_col or blank(v.get('Н5')):
            continue                                                   # легенда, примечания
        title = text(v['Т1']) if not blank(v.get('Т1')) else (cur['title'] if cur else '')
        if cur is None or cur['title'] != title:
            cur = {'title': title, 'rows': []}
            scenarios.append(cur)

        step = text(v['Н5'])
        status = step_status(step, not blank(v.get('О5')))
        f = {'СтатусСтроки': status}
        for code, field in FIELDS.items():
            val = v.get(code)
            if blank(val):
                continue
            if isinstance(val, str) and val.startswith('='):
                val = evaluate(val, cells)
            if field in NUMBER_FIELDS:
                val = number(val)
            elif isinstance(val, (datetime.date, datetime.datetime)):
                val = date(val)
            else:
                val = text(val)
            f[field] = val
        if 'ФормаОплаты' in f:
            f['ФормаОплаты'] = PAYMENT_FORMS.get(f['ФормаОплаты'].lower(), f['ФормаОплаты'])
        if 'Оплачено' in f:
            f['Оплачено'] = PAID.get(f['Оплачено'].lower(), f['Оплачено'])
        if 'ВидУслуги' in f:
            f['ВидУслуги'] = f['ВидУслуги'][:1].upper() + f['ВидУслуги'][1:]
        if status in AFTER_ASSEMBLY and 'ДатаЗаказа' in f:
            f['ДатаОтправкиНаСборку'] = max(f['ДатаЗаказа'], f.get('ДатаПоступления', ''))

        n = len(cur['rows']) + 1
        note = text(v['Т2']).replace('\n', ' ').strip(' /') if not blank(v.get('Т2')) else ''
        f['КомментарийКСтроке2'] = 'Шаг %d «%s»%s' % (n, step, ' — ' + note if note else '')
        f['КомментарийКСтроке'] = title

        client = {fld: text(v[code]) for code, fld in CLIENT.items() if not blank(v.get(code))}
        client.setdefault('НоменклатураКлиента', step)                # шаг без детали (выдача авто)
        row = {'excelRow': r, 'step': step, 'client': client, 'fields': f}
        for key in ('В5', 'В6', 'В7'):
            if not blank(v.get(key)):
                cur.setdefault(key, text(v[key]))
        if 'ДатаЗаказа' in f:
            cur.setdefault('date', f['ДатаЗаказа'])
        cur['rows'].append(row)

        if 'Номенклатура' in f:
            unit = text(v.get('М4', ''))
            nomenclature.setdefault(f['Номенклатура'], {
                'name': f['Номенклатура'], 'article': '' if blank(v.get('М2')) else text(v['М2']),
                'code': '' if blank(v.get('М3')) else text(v['М3']), 'unit': unit,
                'type': 'Работа' if unit.lower() in ('н/ч', 'ч') else 'Запас'})
        if not blank(v.get('В6')):
            c = counterparties.setdefault(text(v['В6']), [])
            if not blank(v.get('В7')) and text(v['В7']) not in c:
                c.append(text(v['В7']))
        if 'Поставщик' in f:
            counterparties.setdefault(f['Поставщик'], [])
        for lst, key in ((service_types, 'ВидУслуги'), (payment_forms, 'ФормаОплаты')):
            if key in f and f[key] not in lst:
                lst.append(f[key])

    orders = []
    for i, s in enumerate(scenarios, 1):
        for row in s['rows']:
            if row['fields'].get('Договор') is None and s.get('В7'):
                row['fields']['Договор'] = s['В7']
            if 'ДатаЗаказа' not in row['fields'] and s.get('date'):     # шаг без В3 (выдача авто) — дата заказа сценария
                row['fields']['ДатаЗаказа'] = s['date']
        orders.append({'clientNumber': 'SEED-%s-%d' % (s.get('В5', '0'), i), 'title': s['title'],
                       'customer': s.get('В6'), 'contract': s.get('В7'), 'date': s.get('date'),
                       'rows': s['rows']})
    return {
        'nomenclature': list(nomenclature.values()),
        'counterparties': [{'name': k, 'contracts': c} for k, c in counterparties.items()],
        'serviceTypes': service_types,
        'paymentForms': payment_forms,
        'orders': orders,
    }


STATIC = {
    '_comment': 'Демо-данные АРМ 3.0 (схема регистра v2.14+): путь детали по процессам из seed/Шаблон деталей.xlsx. '
                'Сгенерировано tools/seed_from_xlsx.py — правьте таблицу или генератор, не этот файл. '
                'Заказ покупателя на сценарий (номер клиента SEED-*), строка АРМ на каждый шаг; поля строки — '
                'по полям регистра Арм_ДанныеЗакупокИПродаж (значения ссылок — наименования).',
    'minExtensionFields': ['ДатаОтправкиНаСборку', 'ЯчейкаСклада', 'НомерСчета', 'ДатаСчета', 'КомментарийКСтроке2',
                           'ДатаЗаказаПоставщику', 'СрокПоставки', 'НомерУПД', 'ДатаУПД', 'КоличествоВПути'],
    'unitMap': {'литр': 'л (дм3)', 'л': 'л (дм3)', 'н/ч': 'ч', 'шт': 'шт', 'кг': 'кг'},
    'adminUser': 'Админ',
    'roles': [
        {'matrix': 'manager', 'arm': 'Менеджер', 'user': 'Менеджер (seed)'},
        {'matrix': 'storekeeper', 'arm': 'Кладовщик', 'user': 'Кладовщик (seed)'},
        {'matrix': 'chief_mechanic', 'arm': 'ГлавныйМеханик', 'user': 'Главный механик (seed)'},
        {'matrix': 'admin', 'arm': 'Администратор', 'user': 'Администратор (seed)'},
        {'matrix': 'supplier', 'arm': 'Поставщик', 'user': 'Поставщик (seed)'},
        {'matrix': 'client', 'arm': 'Клиент', 'user': 'Клиент (seed)'},
        {'matrix': None, 'arm': 'Снабжение', 'user': 'Снабжение (seed)'},
    ],
}


def main(argv):
    xlsx = argv[1] if len(argv) > 1 else XLSX
    out = argv[2] if len(argv) > 2 else OUT
    data = dict(STATIC)
    data['source'] = os.path.relpath(xlsx, ROOT).replace('\\', '/')
    data.update(build(xlsx))
    io.open(out, 'w', encoding='utf-8', newline='\n').write(json.dumps(data, ensure_ascii=False, indent=1) + '\n')
    rows = sum(len(o['rows']) for o in data['orders'])
    print('%s: сценариев %d, строк %d, номенклатуры %d, контрагентов %d' % (
        out, len(data['orders']), rows, len(data['nomenclature']), len(data['counterparties'])))
    for o in data['orders']:
        print('  %s «%s»: %s' % (o['clientNumber'], o['title'], ' → '.join(r['fields']['СтатусСтроки'] for r in o['rows'])))


if __name__ == '__main__':
    main(sys.argv)
